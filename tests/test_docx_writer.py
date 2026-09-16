from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE

from docx_writer import build_docx_bytes
from policy import load_policy
from report_builder import build_report_context
from test_report_builder import make_annual, make_state_use


def test_word_output_uses_canonical_content_without_source_provenance():
    policy = load_policy()
    ctx = build_report_context(make_annual(), make_state_use(), "Most", "August 2026", policy)
    raw = build_docx_bytes(ctx)
    doc = Document(BytesIO(raw))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Lower Basin Projected Water Use: 2026 - 2028" in text
    assert "Arizona: 2.049 maf" in text
    assert "Reduced delivery of 250 kaf" in text
    assert "CRMMS_TEST_MOST" not in text
    assert "report_policy.json" not in text
    assert "Source" not in text


def test_word_package_has_no_custom_or_timestamp_provenance():
    policy = load_policy()
    ctx = build_report_context(make_annual(), make_state_use(), "Most", "August 2026", policy)
    raw = build_docx_bytes(ctx)
    with ZipFile(BytesIO(raw)) as package:
        names = set(package.namelist())
        assert "docProps/custom.xml" not in names
        core = package.read("docProps/core.xml").decode("utf-8")
    lowered = core.lower()
    assert "crmms_test_most" not in lowered
    assert "projected state use-test" not in lowered
    assert "report_policy" not in lowered
    assert "dcterms:created" not in lowered
    assert "dcterms:modified" not in lowered


def test_word_tables_use_quarter_inch_centered_rows_and_no_forced_break_before_conservation():
    policy = load_policy()
    ctx = build_report_context(make_annual(), make_state_use(), "Most", "August 2026", policy)
    raw = build_docx_bytes(ctx)
    doc = Document(BytesIO(raw))

    assert len(doc.tables) == 2
    for table in doc.tables:
        for row in table.rows:
            assert row.height is not None
            assert abs(row.height.inches - 0.25) < 0.001
            assert row.height_rule == WD_ROW_HEIGHT_RULE.EXACTLY
            for cell in row.cells:
                assert cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
                assert all(p.alignment == 1 for p in cell.paragraphs)

    body_xml = doc.element.body.xml
    assert 'w:type="page"' not in body_xml


def test_word_notes_do_not_include_seis_rod_note_and_min_subtitle_has_no_powell_line():
    policy = load_policy()
    min_ctx = build_report_context(make_annual(), make_state_use("Min"), "Min", "August 2026", policy)
    raw = build_docx_bytes(min_ctx)
    doc = Document(BytesIO(raw))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Probable Minimum 24-Month Study" in text
    assert "For the 6 maf & 7 maf Powell Release Scenarios" not in text
    assert "SEIS ROD" not in text
