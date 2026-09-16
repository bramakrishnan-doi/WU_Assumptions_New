"""Render a canonical ReportContext to Word using the curated reference report.

The reference document remains the authoritative source for page setup, named
styles, fonts and numbering definitions. No report business logic lives here.
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor

from report_models import ReportContext

BASE_DIR = Path(__file__).resolve().parent
REFERENCE_DOCX_PATH = BASE_DIR / "reference_report.docx"
WHITE_FILL = "FFFFFF"
TABLE_FONT = "Arial"
STATE_MAF_BLUE = RGBColor(0x00, 0x70, 0xC0)


def _load_base_document() -> DocumentType:
    if not REFERENCE_DOCX_PATH.exists():
        raise FileNotFoundError(
            f"Reference report not found at {REFERENCE_DOCX_PATH}. "
            "The application requires reference_report.docx beside app.py."
        )
    return Document(str(REFERENCE_DOCX_PATH))


def _find_report_bullet_num_id(doc: DocumentType) -> int:
    """Find a numbering definition with the report's bullet/dash hierarchy.

    This avoids hard-coding a numId, which can change when the template is
    edited in Word. We identify the abstract definition by glyphs, then find
    any concrete numId that points to it.
    """
    root = doc.part.numbering_part.element
    wanted_abstract = None
    for abstract in root.findall(qn("w:abstractNum")):
        glyphs = {}
        for level in abstract.findall(qn("w:lvl")):
            ilvl = level.get(qn("w:ilvl"))
            text = level.find(qn("w:lvlText"))
            fmt = level.find(qn("w:numFmt"))
            if text is not None and fmt is not None and fmt.get(qn("w:val")) == "bullet":
                glyphs[ilvl] = text.get(qn("w:val"))
        if glyphs.get("0") in {"•", "\uf0b7"} and glyphs.get("1") in {"–", "-", "—"}:
            wanted_abstract = abstract.get(qn("w:abstractNumId"))
            break
    if wanted_abstract is None:
        raise RuntimeError(
            "reference_report.docx does not contain the expected bullet/dash numbering definition."
        )
    for num in root.findall(qn("w:num")):
        abstract_id = num.find(qn("w:abstractNumId"))
        if abstract_id is not None and abstract_id.get(qn("w:val")) == wanted_abstract:
            return int(num.get(qn("w:numId")))
    raise RuntimeError("Could not find a usable bullet numbering instance in reference_report.docx.")




def _clean_document_metadata(doc: DocumentType, context: ReportContext) -> None:
    """Remove stale template metadata without adding generation provenance."""
    props = doc.core_properties
    props.title = f"Lower Basin Projected Water Use: {context.title_years}"
    props.subject = ""
    props.author = ""
    props.last_modified_by = ""
    props.keywords = ""
    props.comments = ""
    props.category = ""
    props.revision = 1

    # python-docx does not expose a way to set created/modified to None.
    # Remove those XML elements so the generated DOCX contains no generation
    # timestamp and does not retain stale dates from the reference document.
    core = props._element
    for local_name in ("created", "modified"):
        element = core.find(f"{{http://purl.org/dc/terms/}}{local_name}")
        if element is not None:
            core.remove(element)


def _clear_body(doc: DocumentType) -> None:
    body = doc.element.body
    sect_pr = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def _style_exists(doc: DocumentType, name: str) -> bool:
    try:
        doc.styles[name]
        return True
    except KeyError:
        return False


def _safe_style(doc: DocumentType, name: str, fallback: str = "Normal") -> str:
    return name if _style_exists(doc, name) else fallback


def _add_bullet_numbering(paragraph, level: int, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    old = p_pr.find(qn("w:numPr"))
    if old is not None:
        p_pr.remove(old)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(1 if level else 0))
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    p_pr.append(num_pr)


def _add_bullet_line(doc: DocumentType, text: str, num_id: int, level: int = 0) -> None:
    paragraph = doc.add_paragraph(style=_safe_style(doc, "Compact" if level else "Normal"))
    _add_bullet_numbering(paragraph, level, num_id)
    paragraph.add_run(text)


def _add_heading(doc: DocumentType, text: str, level: int) -> None:
    heading = doc.add_heading(level=level)
    heading.add_run(text)


def _add_state_heading(doc: DocumentType, state_name: str, maf_value: str, level: int) -> None:
    heading = doc.add_heading(level=level)
    heading.add_run(f"{state_name}: ")
    value = heading.add_run(f"{maf_value} maf")
    value.font.color.rgb = STATE_MAF_BLUE


def _add_year_section(doc: DocumentType, section, num_id: int) -> None:
    _add_heading(doc, str(section.year), level=3)
    p = doc.add_paragraph(style=_safe_style(doc, "First Paragraph"))
    p.add_run("Total projected water use ")
    p.add_run(f"({section.total_use_maf} maf)").bold = True
    p.add_run(" – ")
    disclaimer = p.add_run(section.disclaimer_text)
    disclaimer.italic = True
    disclaimer.font.highlight_color = WD_COLOR_INDEX.YELLOW

    _add_heading(doc, f"U.S. Contractors: {section.us_contractors_maf} maf", level=4)
    for state in section.states:
        _add_state_heading(doc, state.heading, state.maf_value, level=5)
        for bullet in state.bullets:
            _add_bullet_line(doc, bullet.text, num_id, bullet.level)

    _add_heading(doc, section.mexico_heading, level=4)
    for bullet in section.mexico_bullets:
        _add_bullet_line(doc, bullet.text, num_id, bullet.level)


def _set_cell_background(cell, hex_color: str = WHITE_FILL) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:shd"))
    if old is not None:
        tc_pr.remove(old)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_color)
    tc_pr.append(shading)


def _set_cell_borders(cell, size: str = "8", color: str = "000000") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:tcBorders"))
    if old is not None:
        tc_pr.remove(old)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tc_pr.append(borders)


def _style_table_cell_text(cell, bold: bool) -> None:
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if not paragraph.runs:
            paragraph.add_run("")
        for run in paragraph.runs:
            run.font.name = TABLE_FONT
            run.font.bold = bold
            run.font.color.rgb = RGBColor(0, 0, 0)


def _format_table_cell(cell, text: str, bold: bool = False) -> None:
    cell.text = text
    _set_cell_background(cell)
    _set_cell_borders(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _style_table_cell_text(cell, bold)


def _set_report_table_row_height(row) -> None:
    row.height = Inches(0.25)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY


def _add_ics_table(doc: DocumentType, table_data) -> None:
    _add_heading(doc, "ICS Totals", level=3)
    p1 = doc.add_paragraph(style=_safe_style(doc, "Compact"))
    p1.add_run("Projected ICS Total Storage at the end of CY ")
    p1.add_run(f"{table_data.final_year_label}: ")
    p1.add_run(f"{table_data.final_year_total_maf} maf").bold = True
    doc.add_paragraph(
        "Projected ICS Storage Balances at the end of each calendar year in the study are as follows:",
        style=_safe_style(doc, "Compact"),
    )

    table = doc.add_table(rows=1, cols=1 + len(table_data.year_labels))
    if _style_exists(doc, "Normal Table"):
        table.style = "Normal Table"
    headers = ["State (volumes in AF)", *table_data.year_labels]
    for cell, text in zip(table.rows[0].cells, headers):
        _format_table_cell(cell, text, bold=True)
    _set_report_table_row_height(table.rows[0])
    for row in table_data.rows:
        cells = table.add_row().cells
        values = [row.state, *row.values]
        for cell, text in zip(cells, values):
            _format_table_cell(cell, text, bold=row.is_total)
        _set_report_table_row_height(table.rows[-1])


def _add_conservation_table(doc: DocumentType, summary) -> None:
    _add_heading(doc, "Modeled Lower Basin Conservation Actions", level=3)
    table = doc.add_table(rows=1, cols=2 + len(summary.year_labels))
    if _style_exists(doc, "Normal Table"):
        table.style = "Normal Table"
    headers = ["State", *summary.year_labels, "Total"]
    for cell, text in zip(table.rows[0].cells, headers):
        _format_table_cell(cell, text, bold=True)
    _set_report_table_row_height(table.rows[0])
    for row in summary.rows:
        cells = table.add_row().cells
        values = [row.state, *row.values, row.total]
        for cell, text in zip(cells, values):
            _format_table_cell(cell, text, bold=row.bold)
        _set_report_table_row_height(table.rows[-1])


def _add_notes(doc: DocumentType, context: ReportContext, num_id: int) -> None:
    doc.add_paragraph()
    _add_heading(doc, "Notes and Disclaimers", level=2)
    # Reference report uses top-level bullets here; keep that same list hierarchy.
    _add_bullet_line(doc, context.conservation_disclaimer, num_id, level=0)


def build_docx(context: ReportContext) -> DocumentType:
    """Build a Word Document from the canonical context."""
    doc = _load_base_document()
    num_id = _find_report_bullet_num_id(doc)
    _clear_body(doc)
    _clean_document_metadata(doc, context)

    title = doc.add_paragraph(style=_safe_style(doc, "Title"))
    title.add_run(f"Lower Basin Projected Water Use: {context.title_years}")

    subtitle = doc.add_paragraph(style=_safe_style(doc, "Subtitle"))
    subtitle.add_run(
        f"As Modeled in the {context.mon_year}: {context.scenario_label} 24-Month Study"
    )
    if context.powell_release_subtitle:
        subtitle.add_run().add_break()
        subtitle.add_run(context.powell_release_subtitle)

    for year_section in context.years:
        _add_year_section(doc, year_section, num_id)
    _add_ics_table(doc, context.ics_table)
    if context.show_conservation_summary:
        _add_conservation_table(doc, context.conservation_summary)
    _add_notes(doc, context, num_id)
    return doc


def build_docx_bytes(context: ReportContext) -> bytes:
    out = BytesIO()
    build_docx(context).save(out)
    return out.getvalue()
