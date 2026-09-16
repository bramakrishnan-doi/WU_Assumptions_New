from openpyxl import Workbook

from excel_source import _load_year_sheet
from policy import load_policy


def test_scenario_specific_columns_are_used():
    policy = load_policy()
    wb = Workbook()
    ws = wb.active
    ws.title = "2026"

    # Same logical row, deliberately different scenario values.
    ws["A1"] = "Most Probable"
    ws["F1"] = "Probable Minimum"
    ws["K1"] = "Probable Maximum"
    ws["A11"] = "California- Water left in Mead"
    ws["B11"] = -185_889
    ws["D11"] = "Most note"
    ws["F11"] = "California- Water left in Mead"
    ws["G11"] = -157_889
    ws["I11"] = "Min note"
    ws["K11"] = "California- Water left in Mead"
    ws["L11"] = -199_674
    ws["N11"] = "Max note"

    most = _load_year_sheet(ws, "Most", policy).find_entry("California- Water left in Mead")
    minimum = _load_year_sheet(ws, "Min", policy).find_entry("California- Water left in Mead")
    maximum = _load_year_sheet(ws, "Max", policy).find_entry("California- Water left in Mead")

    assert (most.value, most.value_cell, most.note) == (-185_889, "B11", "Most note")
    assert (minimum.value, minimum.value_cell, minimum.note) == (-157_889, "G11", "Min note")
    assert (maximum.value, maximum.value_cell, maximum.note) == (-199_674, "L11", "Max note")
