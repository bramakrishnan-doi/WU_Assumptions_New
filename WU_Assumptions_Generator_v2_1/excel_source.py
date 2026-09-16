"""Projected State Use workbook reader with scenario-aware columns and source tracing."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Sequence

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from openpyxl.utils.cell import range_boundaries

from policy import ReportPolicy


class ExcelValidationError(Exception):
    def __init__(self, errors: List[str]):
        super().__init__("; ".join(errors))
        self.errors = errors


@dataclass(frozen=True)
class ExcelValue:
    value: Optional[float]
    sheet: str
    label: str
    value_cell: str
    note: Optional[str] = None
    note_cell: Optional[str] = None


@dataclass
class YearSheetData:
    sheet_name: str
    scenario: str
    entries: List[ExcelValue] = field(default_factory=list)

    def find_entry(self, needle: str) -> Optional[ExcelValue]:
        needle_norm = _normalize_for_search(needle)
        for entry in self.entries:
            if needle_norm in _normalize_for_search(entry.label):
                return entry
        return None

    def find(self, needle: str) -> Optional[float]:
        entry = self.find_entry(needle)
        return entry.value if entry else None

    def find_note(self, needle: str) -> Optional[str]:
        entry = self.find_entry(needle)
        return entry.note if entry else None


@dataclass
class ConservationRow:
    number: Optional[float]
    contractor: str
    values_by_year: Dict[int, float]
    cells_by_year: Dict[int, str]
    sheet_name: str


@dataclass
class ConservationTables:
    ca: List[ConservationRow]
    non_cawcd: List[ConservationRow]
    non_cawcd_all: List[ConservationRow]
    cawcd: List[ConservationRow]
    year_columns: List[int]


@dataclass
class SummaryRow:
    state: str
    values_by_year: Dict[int, float]
    cells_by_year: Dict[int, str]
    total: Optional[float]
    total_cell: Optional[str]
    sheet_name: str


@dataclass
class StateUseData:
    workbook_filename: str
    scenario: str
    year_sheets: Dict[int, YearSheetData]
    conservation: ConservationTables
    summary_rows: List[SummaryRow]
    has_year3_sheet: bool
    warnings: List[str] = field(default_factory=list)


def _to_float(value) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        return None


def _normalize_for_search(text: str) -> str:
    normalized = str(text).lower()
    normalized = normalized.replace("–", "-").replace("—", "-")
    normalized = re.sub(r"\s*-\s*", "-", normalized)
    normalized = re.sub(r"\s+", " ", normalized).strip()
    return normalized


def _load_year_sheet(ws, scenario: str, policy: ReportPolicy) -> YearSheetData:
    scenario_cfg = policy.scenario(scenario)
    label_col = column_index_from_string(scenario_cfg.label_column)
    value_col = column_index_from_string(scenario_cfg.value_column)
    notes_col = column_index_from_string(scenario_cfg.notes_column)

    data = YearSheetData(sheet_name=ws.title, scenario=scenario)
    for row_idx in range(1, ws.max_row + 1):
        label = ws.cell(row_idx, label_col).value
        if label is None or not str(label).strip():
            continue
        raw_value = ws.cell(row_idx, value_col).value
        note = ws.cell(row_idx, notes_col).value
        data.entries.append(
            ExcelValue(
                value=_to_float(raw_value),
                sheet=ws.title,
                label=str(label).strip(),
                value_cell=ws.cell(row_idx, value_col).coordinate,
                note=str(note).strip() if note is not None and str(note).strip() else None,
                note_cell=ws.cell(row_idx, notes_col).coordinate,
            )
        )
    return data


def _load_conservation_block(ws, range_text: str) -> tuple[List[ConservationRow], List[int]]:
    min_col, configured_row, max_col, max_row = range_boundaries(range_text)
    contractor_col = min_col + 1

    # The configured range defines the block's columns and expected maximum
    # height, but the header row is discovered by text. This tolerates rows
    # being inserted above the table in the monthly workbook.
    header_row = None
    search_limit = min(ws.max_row, max(configured_row + 60, 120))
    for row_idx in range(1, search_limit + 1):
        if _normalize_for_search(ws.cell(row_idx, contractor_col).value or "") == "contractor":
            header_row = row_idx
            break
    if header_row is None:
        col_letter = ws.cell(1, contractor_col).column_letter
        raise ExcelValidationError([
            f"Could not locate a 'Contractor' header in column {col_letter} on '{ws.title}' "
            f"for configured conservation block {range_text}."
        ])

    year_columns: List[int] = []
    year_col_indexes: List[int] = []
    for col_idx in range(contractor_col + 1, max_col + 1):
        header = ws.cell(header_row, col_idx).value
        digits = "".join(ch for ch in str(header or "") if ch.isdigit())
        if len(digits) == 4:
            year_columns.append(int(digits))
            year_col_indexes.append(col_idx)

    if not year_columns:
        raise ExcelValidationError([
            f"No calendar-year columns were found in {ws.title}!{range_text}."
        ])

    rows: List[ConservationRow] = []
    expected_height = max_row - configured_row
    data_end_row = min(ws.max_row, header_row + expected_height)
    for row_idx in range(header_row + 1, data_end_row + 1):
        contractor = ws.cell(row_idx, contractor_col).value
        if contractor is None or not str(contractor).strip():
            continue
        values_by_year: Dict[int, float] = {}
        cells_by_year: Dict[int, str] = {}
        for year, col_idx in zip(year_columns, year_col_indexes):
            cell = ws.cell(row_idx, col_idx)
            value = _to_float(cell.value)
            if value is not None:
                values_by_year[year] = value
                cells_by_year[year] = cell.coordinate
        rows.append(
            ConservationRow(
                number=_to_float(ws.cell(row_idx, min_col).value),
                contractor=str(contractor).strip(),
                values_by_year=values_by_year,
                cells_by_year=cells_by_year,
                sheet_name=ws.title,
            )
        )
    return rows, year_columns


def _load_conservation_tables(wb, policy: ReportPolicy) -> ConservationTables:
    sheet_name = str(policy.excel["system_conservation_sheet"])
    if sheet_name not in wb.sheetnames:
        raise ExcelValidationError([f"Required sheet '{sheet_name}' was not found in the workbook."])
    ws = wb[sheet_name]
    blocks = policy.excel["conservation_blocks"]

    cawcd, years_a = _load_conservation_block(ws, str(blocks["cawcd"]["range"]))
    non_cawcd_all, years_b = _load_conservation_block(ws, str(blocks["non_cawcd"]["range"]))
    ca, years_c = _load_conservation_block(ws, str(blocks["california"]["range"]))
    if not (years_a == years_b == years_c):
        raise ExcelValidationError([
            "Configured System Conservation blocks do not contain the same calendar-year columns."
        ])

    non_cawcd = [r for r in non_cawcd_all if "242 wellfield" not in r.contractor.lower()]
    return ConservationTables(
        ca=ca,
        non_cawcd=non_cawcd,
        non_cawcd_all=non_cawcd_all,
        cawcd=cawcd,
        year_columns=years_a,
    )


def _find_summary_header(ws, start_year: int, label: str) -> tuple[int, int, List[int]] | None:
    wanted = _normalize_for_search(label)
    max_search_row = min(ws.max_row, 250)
    max_search_col = min(ws.max_column, 40)
    for row_idx in range(1, max_search_row + 1):
        for col_idx in range(1, max_search_col + 1):
            if _normalize_for_search(ws.cell(row_idx, col_idx).value or "") != wanted:
                continue
            years: List[int] = []
            scan_col = col_idx + 1
            while scan_col <= min(ws.max_column, col_idx + 8):
                raw = ws.cell(row_idx, scan_col).value
                try:
                    year = int(raw)
                except (TypeError, ValueError):
                    break
                if 1900 <= year <= 2200:
                    years.append(year)
                    scan_col += 1
                    continue
                break
            if years and years[0] == start_year:
                return row_idx, col_idx, years
    return None


def _load_summary_table(wb, start_year: int, policy: ReportPolicy) -> List[SummaryRow]:
    sheet_name = str(policy.excel["summary_sheet"])
    if sheet_name not in wb.sheetnames:
        raise ExcelValidationError([f"Required sheet '{sheet_name}' was not found in the workbook."])
    ws = wb[sheet_name]
    header_info = _find_summary_header(ws, start_year, str(policy.excel.get("summary_header_label", "State")))
    if header_info is None:
        raise ExcelValidationError([
            f"Could not locate the conservation summary header containing year {start_year} on '{sheet_name}'."
        ])
    header_row, state_col, years = header_info
    total_col = state_col + 1 + len(years)
    expected_rows = [str(v) for v in policy.excel.get("summary_expected_rows", [])]
    expected_norm = {_normalize_for_search(v) for v in expected_rows}

    rows: List[SummaryRow] = []
    for row_idx in range(header_row + 1, min(ws.max_row, header_row + 12) + 1):
        state = ws.cell(row_idx, state_col).value
        if state is None or not str(state).strip():
            if rows:
                break
            continue
        state_str = str(state).strip()
        values_by_year: Dict[int, float] = {}
        cells_by_year: Dict[int, str] = {}
        for offset, year in enumerate(years, start=1):
            cell = ws.cell(row_idx, state_col + offset)
            value = _to_float(cell.value)
            if value is not None:
                values_by_year[year] = value
                cells_by_year[year] = cell.coordinate
        total_cell_obj = ws.cell(row_idx, total_col)
        rows.append(
            SummaryRow(
                state=state_str,
                values_by_year=values_by_year,
                cells_by_year=cells_by_year,
                total=_to_float(total_cell_obj.value),
                total_cell=total_cell_obj.coordinate,
                sheet_name=ws.title,
            )
        )
        if expected_norm and expected_norm.issubset({_normalize_for_search(r.state) for r in rows}):
            break

    found_norm = {_normalize_for_search(r.state) for r in rows}
    missing = [v for v in expected_rows if _normalize_for_search(v) not in found_norm]
    if missing:
        raise ExcelValidationError([
            f"Conservation summary table on '{sheet_name}' is missing expected row(s): {', '.join(missing)}."
        ])
    return rows


def _scenario_header_warning(year_data: YearSheetData, policy: ReportPolicy) -> Optional[str]:
    display = policy.scenario(year_data.scenario).display_name.lower()
    first_labels = " ".join(e.label.lower() for e in year_data.entries[:6])
    scenario_tokens = {
        "Most": ("most probable", "most"),
        "Min": ("probable minimum", "minimum", "min"),
        "Max": ("probable maximum", "maximum", "max"),
    }
    if not any(token in first_labels for token in scenario_tokens.get(year_data.scenario, ())):
        return (
            f"Sheet '{year_data.sheet_name}' did not contain an obvious {display} scenario heading "
            f"in the configured {policy.scenario(year_data.scenario).label_column} column. Values were still read from the configured scenario block."
        )
    return None


def load_state_use_workbook(
    data: bytes,
    filename: str,
    start_year: int,
    scenarios: Sequence[str],
    policy: ReportPolicy,
) -> Dict[str, StateUseData]:
    """Load the workbook once and return scenario-specific views.

    The System Conservation and summary tables are shared, while each calendar
    year sheet is read from the scenario-specific A:D, F:I, or K:N block
    defined by policy.
    """
    if not filename.lower().endswith(".xlsx"):
        raise ExcelValidationError(["The Projected State Use source must be an .xlsx workbook."])
    scenarios = list(dict.fromkeys(scenarios))
    unknown = [s for s in scenarios if s not in policy.scenario_keys]
    if unknown:
        raise ExcelValidationError([f"Unknown scenario(s): {', '.join(unknown)}."])
    if not scenarios:
        raise ExcelValidationError(["At least one scenario is required."])

    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExcelValidationError([f"Could not open the Excel workbook: {exc}"]) from exc

    common_warnings: List[str] = []
    if getattr(wb.calculation, "calcMode", None) == "manual":
        common_warnings.append(
            "Workbook calculation mode is Manual. Open it in Excel, recalculate, and save before generating reports so cached formula values are current."
        )

    # Validate required year sheets once.
    missing_required_sheets = [
        str(start_year + offset)
        for offset in range(min(2, policy.target_year_count))
        if str(start_year + offset) not in wb.sheetnames
    ]
    if missing_required_sheets:
        raise ExcelValidationError([
            "Required year sheet(s) not found: " + ", ".join(missing_required_sheets)
        ])

    conservation = _load_conservation_tables(wb, policy)
    summary_rows = _load_summary_table(wb, start_year, policy)
    has_year3_sheet = str(start_year + 2) in wb.sheetnames
    report_years = [start_year, start_year + 1] + ([start_year + 2] if has_year3_sheet else [])

    errors: List[str] = []
    missing_cons_years = [y for y in report_years if y not in conservation.year_columns]
    if missing_cons_years:
        errors.append(
            "System Conservation table is missing report year(s): " + ", ".join(map(str, missing_cons_years))
        )

    all_conservation_rows = [*conservation.ca, *conservation.non_cawcd_all, *conservation.cawcd]
    for year in report_years:
        numeric_count = sum(1 for row in all_conservation_rows if year in row.values_by_year)
        if numeric_count == 0:
            errors.append(
                f"System Conservation contains no cached numeric values for {year}. "
                "Open the workbook in Excel, recalculate, and save it."
            )

    results: Dict[str, StateUseData] = {}
    for scenario in scenarios:
        warnings = list(common_warnings)
        year_sheets: Dict[int, YearSheetData] = {}
        for offset in range(policy.target_year_count):
            year = start_year + offset
            sheet_name = str(year)
            if sheet_name not in wb.sheetnames:
                continue
            data_for_year = _load_year_sheet(wb[sheet_name], scenario, policy)
            year_sheets[year] = data_for_year
            warning = _scenario_header_warning(data_for_year, policy)
            if warning:
                warnings.append(warning)
            for required_label in policy.excel.get("year_sheet_required_labels", []):
                if data_for_year.find_entry(str(required_label)) is None:
                    errors.append(
                        f"Sheet '{sheet_name}' scenario '{scenario}' is missing expected label containing '{required_label}'."
                    )

        y1 = year_sheets.get(start_year)
        y1_mexico = y1.find_entry("Mexico Use") if y1 else None
        if y1_mexico is None or y1_mexico.value is None:
            errors.append(
                f"{start_year} scenario '{scenario}' has no numeric cached value for Mexico Use. "
                "Open the workbook in Excel, recalculate, and save it."
            )

        if policy.shows_conservation_summary(scenario):
            summary_years = {y for row in summary_rows for y in row.values_by_year}
            missing_summary_years = [y for y in report_years if y not in summary_years]
            if missing_summary_years:
                errors.append(
                    "Conservation summary table has no cached values for report year(s): "
                    + ", ".join(map(str, missing_summary_years))
                    + ". Open the workbook in Excel, recalculate, and save it."
                )

        results[scenario] = StateUseData(
            workbook_filename=filename,
            scenario=scenario,
            year_sheets=year_sheets,
            conservation=conservation,
            summary_rows=summary_rows,
            has_year3_sheet=has_year3_sheet,
            warnings=warnings,
        )

    if errors:
        raise ExcelValidationError(errors[:50])
    return results


def load_state_use_excel(
    data: bytes,
    filename: str,
    start_year: int,
    scenario: str,
    policy: ReportPolicy,
) -> StateUseData:
    return load_state_use_workbook(data, filename, start_year, [scenario], policy)[scenario]
